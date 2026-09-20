"""
Reads the latest monthly P&L workbook from data-raw/ and exports the
"P&L - Monthly Snapshot" tab to frontend/src/data/report.json for the portal.

Usage:
    python export_report.py

The workbook must already have cached formula values (i.e. it was saved by
Excel, or run through the xlsx skill's recalc.py) — openpyxl cannot evaluate
formulas itself.
"""
import datetime
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "data-raw"
OUT_PATH = ROOT / "frontend" / "src" / "data" / "report.json"
SHEET_NAME = "P&L - Monthly Snapshot"
HEADER_ROW = 10


def latest_workbook() -> Path:
    files = sorted(SRC_DIR.glob("*.xlsx"))
    if not files:
        raise SystemExit(f"No .xlsx files found in {SRC_DIR}. Commit a monthly workbook there first.")
    return files[-1]


def read_blocks(ws, header_row: int):
    """Find each 'Line Item' header in the row and the metric columns that follow it."""
    label_cols = [c for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value == "Line Item"]
    blocks = []
    for i, col in enumerate(label_cols):
        end_col = (label_cols[i + 1] - 1) if i + 1 < len(label_cols) else ws.max_column
        headers = []
        for c in range(col + 1, end_col + 1):
            v = ws.cell(header_row, c).value
            if v is None:
                break
            headers.append((c, str(v)))
        blocks.append({"label_col": col, "headers": headers})
    return blocks


def read_table(ws, block, start_row: int):
    rows = []
    r = start_row
    while True:
        label = ws.cell(r, block["label_col"]).value
        if label is None or str(label).strip() == "":
            break
        row = {"item": str(label).strip()}
        for c, h in block["headers"]:
            row[h] = ws.cell(r, c).value
        rows.append(row)
        r += 1
    return rows


def main():
    wb_path = latest_workbook()
    wb = openpyxl.load_workbook(wb_path, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Sheet '{SHEET_NAME}' not found in {wb_path.name}")
    ws = wb[SHEET_NAME]

    blocks = read_blocks(ws, HEADER_ROW)
    if len(blocks) < 3:
        raise SystemExit(
            f"Expected 3 'Line Item' tables (MTD / Rolling Actual / Rolling Budget) "
            f"in row {HEADER_ROW}, found {len(blocks)}. Has the template layout changed?"
        )

    mtd_block, actual_block, budget_block = blocks[0], blocks[1], blocks[2]
    reporting_label = ws["C9"].value or ws["C5"].value

    data = {
        "reportingMonth": str(reporting_label) if reporting_label else None,
        "sourceFile": wb_path.name,
        "generatedAt": datetime.datetime.utcnow().isoformat() + "Z",
        "mtd": read_table(ws, mtd_block, HEADER_ROW + 1),
        "rollingActual": read_table(ws, actual_block, HEADER_ROW + 1),
        "rollingBudget": read_table(ws, budget_block, HEADER_ROW + 1),
    }

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {OUT_PATH.relative_to(ROOT)} from {wb_path.name} ({len(data['mtd'])} MTD line items)")


if __name__ == "__main__":
    sys.exit(main())
